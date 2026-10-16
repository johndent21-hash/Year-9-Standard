#!/usr/bin/env python3
"""Builds the Word versions of the Year 9 5.2 Yearly Exam 2026 and its marking guide from build/exam.json.

    python3 build-docx.py --json build/exam.json --out "../Yearly exam 2026"

build.js runs this after it has written exam.json and the figure PNGs (build/figs, 300 dpi). The Word files mirror
the PDF: A4, the same margins, the same front page (logo, NAME and TEACHER boxes, title lines, Instructions and
Sections-and-Marks boxes), and the same section tables (navy title row repeated on every page, light-blue sub-header
rows, Question | Working Out | Mark columns, row heights for working space, rows that never split across pages).
Text stays editable: Unicode for ½ ² ³ π × ÷, real Word tables, and equation objects for stacked fractions.
Fonts are Cambria (tables), Arial (front page), Times New Roman (name boxes) and Calibri (page numbers): the fonts of
the 2024 exam. LibreOffice shows them with Caladea, Liberation Sans/Serif and Carlito, which have the same widths.
"""
import argparse
import copy
import json
import os
import re
from html.parser import HTMLParser

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_ROW_HEIGHT_RULE, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn
from docx.shared import Mm, Pt, RGBColor

NAVY = '0F3A63'
KBLUE = '1D71B8'
TINT = 'DBE7F4'
TINT2 = 'EEF4FB'
RULE = '1F1F1F'
RED = 'C00000'
GREEN = '2E9B47'
SERIF = 'Cambria'
SANS = 'Arial'
TIMES = 'Times New Roman'
PNFONT = 'Calibri'
LETTERS = 'ABCD'

# page geometry (mm), measured from the 2024 exam
PAGE_W, PAGE_H = 210, 297
M_LEFT = M_RIGHT = 13
M_TOP = 25.4
M_BOTTOM = 24.6
TABLE_W = 184
COLS = [6.7, 66, 99.3, 12]           # number | question | working out | mark
CELL_PAD_X, CELL_PAD_Y = 1.6, 1.2
L11 = 11 * 1.172 * 25.4 / 72         # one line of 11 pt Cambria, mm
L10 = 10 * 1.22 * 25.4 / 72


# ---------------------------------------------------------------- low-level helpers
def el(tag, **attrs):
    e = OxmlElement(tag)
    for k, v in attrs.items():
        e.set(qn(k), str(v))
    return e


def twips(mm):
    return int(round(mm * 1440 / 25.4))


# Word is strict about the order of property elements; these are the schema orders.
TCPR = 'cnfStyle tcW gridSpan hMerge vMerge tcBorders shd noWrap tcMar textDirection tcFitText vAlign hideMark'.split()
TBLPR = 'tblStyle tblpPr tblOverlap bidiVisual tblStyleRowBandSize tblStyleColBandSize tblW jc tblCellSpacing tblInd tblBorders shd tblLayout tblCellMar tblLook tblCaption tblDescription'.split()
BORDERS = 'top start left bottom end right insideH insideV tl2br tr2bl'.split()
MARS = 'top start left bottom end right'.split()
PPR = ('pStyle keepNext keepLines pageBreakBefore framePr widowControl numPr suppressLineNumbers pBdr shd tabs suppressAutoHyphens '
       'kinsoku wordWrap overflowPunct topLinePunct autoSpaceDE autoSpaceDN bidi adjustRightInd snapToGrid spacing ind contextualSpacing '
       'mirrorIndents suppressOverlap jc textDirection textAlignment textboxTightWrap outlineLvl divId cnfStyle rPr sectPr pPrChange').split()
RPR = ('rStyle rFonts b bCs i iCs caps smallCaps strike dstrike outline shadow emboss imprint noProof snapToGrid vanish webHidden color '
       'spacing w kern position sz szCs highlight u effect bdr shd fitText vertAlign rtl cs em lang eastAsianLayout specVanish oMath').split()


def put(parent, child, order):
    """Insert child into parent in schema order, replacing any element with the same tag."""
    tag = child.tag
    for old in parent.findall(tag):
        parent.remove(old)
    name = tag.split('}')[1]
    idx = order.index(name)
    for i, existing in enumerate(list(parent)):
        en = existing.tag.split('}')[1] if '}' in existing.tag else existing.tag
        if en in order and order.index(en) > idx:
            existing.addprevious(child)
            return child
    parent.append(child)
    return child


def set_cell_shading(cell, fill):
    put(cell._tc.get_or_add_tcPr(), el('w:shd', **{'w:val': 'clear', 'w:color': 'auto', 'w:fill': fill}), TCPR)


def set_cell_borders(cell, **sides):
    """sides: top/left/bottom/right = (size_eighths, color) or None for nil."""
    tcPr = cell._tc.get_or_add_tcPr()
    b = tcPr.find(qn('w:tcBorders'))
    if b is None:
        b = put(tcPr, el('w:tcBorders'), TCPR)
    for side, v in sides.items():
        if v is None:
            put(b, el(f'w:{side}', **{'w:val': 'nil'}), BORDERS)
        else:
            put(b, el(f'w:{side}', **{'w:val': 'single', 'w:sz': v[0], 'w:space': 0, 'w:color': v[1]}), BORDERS)


def set_cell_margins(cell, top=None, left=None, bottom=None, right=None):
    tcPr = cell._tc.get_or_add_tcPr()
    m = tcPr.find(qn('w:tcMar'))
    if m is None:
        m = put(tcPr, el('w:tcMar'), TCPR)
    for side, v in (('top', top), ('left', left), ('bottom', bottom), ('right', right)):
        if v is not None:
            put(m, el(f'w:{side}', **{'w:w': twips(v), 'w:type': 'dxa'}), MARS)


def set_cell_width(cell, mm):
    cell.width = Mm(mm)


def table_setup(table, widths, borders=True, pad=(CELL_PAD_Y, CELL_PAD_X), indent=0, sz=6):
    """Fixed layout, exact column widths, single dark rules (or none), default cell padding."""
    tbl = table._tbl
    tblPr = tbl.tblPr
    for tag in ('w:tblStyle',):
        for old in tblPr.findall(qn(tag)):
            tblPr.remove(old)
    put(tblPr, el('w:tblW', **{'w:w': twips(sum(widths)), 'w:type': 'dxa'}), TBLPR)
    if indent:
        put(tblPr, el('w:tblInd', **{'w:w': twips(indent), 'w:type': 'dxa'}), TBLPR)
    b = el('w:tblBorders')
    for side in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
        if borders:
            b.append(el(f'w:{side}', **{'w:val': 'single', 'w:sz': sz, 'w:space': 0, 'w:color': RULE}))
        else:
            b.append(el(f'w:{side}', **{'w:val': 'nil'}))
    put(tblPr, b, TBLPR)
    put(tblPr, el('w:tblLayout', **{'w:type': 'fixed'}), TBLPR)
    m = el('w:tblCellMar')
    for side, v in (('top', pad[0]), ('left', pad[1]), ('bottom', pad[0]), ('right', pad[1])):
        m.append(el(f'w:{side}', **{'w:w': twips(v), 'w:type': 'dxa'}))
    put(tblPr, m, TBLPR)
    grid = tbl.tblGrid
    for i, gc in enumerate(grid.findall(qn('w:gridCol'))):
        gc.set(qn('w:w'), str(twips(widths[i])))
    for row in table.rows:  # cells made with the table get the column widths too (later rows take them from the grid)
        for i, tc in enumerate(row._tr.tc_lst):
            tc.width = Mm(widths[i])
    table.autofit = False


def row_props(row, height=None, rule='atLeast', cant_split=True, header=False):
    trPr = row._tr.get_or_add_trPr()
    if cant_split:
        trPr.append(el('w:cantSplit'))
    if header:
        trPr.append(el('w:tblHeader'))
    if height is not None:
        row.height = Mm(height)
        row.height_rule = WD_ROW_HEIGHT_RULE.EXACTLY if rule == 'exact' else WD_ROW_HEIGHT_RULE.AT_LEAST


def merge(row, a, b, widths=None):
    c = row.cells[a].merge(row.cells[b])
    ps = c.paragraphs
    for p in ps[1:]:
        p._p.getparent().remove(p._p)
    c.width = Mm(sum((widths or COLS)[a:b + 1]))
    return c


def font(run, name=SERIF, size=11, bold=None, italic=None, color=None, sup=False, sub=False):
    run.font.name = name
    rPr = run._r.get_or_add_rPr()
    rf = rPr.find(qn('w:rFonts'))
    if rf is None:
        rf = el('w:rFonts')
        rPr.insert(0, rf)
    for k in ('w:ascii', 'w:hAnsi', 'w:cs', 'w:eastAsia'):
        rf.set(qn(k), name)
    run.font.size = Pt(size)
    if bold is not None:
        run.font.bold = bold
    if italic is not None:
        run.font.italic = italic
    if color:
        run.font.color.rgb = RGBColor.from_string(color)
    if sup:
        run.font.superscript = True
    if sub:
        run.font.subscript = True


def para_fmt(p, align=None, before=0, after=0, line=None, line_exact=None, left=None, hanging=None, first=None, keep_next=False, tabs=()):
    pf = p.paragraph_format
    if align:
        pf.alignment = {'c': WD_ALIGN_PARAGRAPH.CENTER, 'l': WD_ALIGN_PARAGRAPH.LEFT, 'r': WD_ALIGN_PARAGRAPH.RIGHT}[align]
    pf.space_before = Mm(max(0, before))
    pf.space_after = Mm(max(0, after))
    if line_exact is not None:
        pf.line_spacing_rule = WD_LINE_SPACING.EXACTLY
        pf.line_spacing = Mm(line_exact)
    elif line is not None:
        pf.line_spacing = line
    else:
        pf.line_spacing = 1.0
    if left is not None:
        pf.left_indent = Mm(left)
    if hanging is not None:
        pf.first_line_indent = Mm(-hanging)
    if first is not None:
        pf.first_line_indent = Mm(first)
    if keep_next:
        pf.keep_with_next = True
    for t in tabs:
        pf.tab_stops.add_tab_stop(Mm(t))


def tiny_para(container, size=1, page_break=False):
    """A paragraph about 0.35 mm high: separates tables and carries page breaks."""
    p = container.add_paragraph()
    para_fmt(p, line_exact=0.35)
    if page_break:
        p.paragraph_format.page_break_before = True
    r = p.add_run('')
    font(r, size=size)
    return p


def picture(run, path, width_mm):
    """An inline picture with no extra spacing round it (LibreOffice otherwise adds Word's 3.2 mm float spacing)."""
    shape = run.add_picture(path, width=Mm(width_mm))
    inl = shape._inline
    for k in ('distT', 'distB', 'distL', 'distR'):
        inl.set(k, '0')
    return shape


def first_para(cell):
    return cell.paragraphs[0]


def next_para(cell, used):
    """The cell's first paragraph if it is still unused, else a new one."""
    if not used[0]:
        used[0] = True
        return cell.paragraphs[0]
    return cell.add_paragraph()


# ---------------------------------------------------------------- inline HTML -> runs (and equation fractions)
M_NS = 'http://schemas.openxmlformats.org/officeDocument/2006/math'
SUPS = str.maketrans('0123456789-+', '⁰¹²³⁴⁵⁶⁷⁸⁹⁻⁺')


def omml_run(text, italic, color, size):
    t = (text.replace('&', '&amp;').replace('<', '&lt;').replace('>', '&gt;'))
    col = f'<w:color w:val="{color}"/>' if color else ''
    return (f'<m:r><m:rPr><m:sty m:val="{"i" if italic else "p"}"/></m:rPr>'
            f'<w:rPr><w:rFonts w:ascii="Cambria Math" w:hAnsi="Cambria Math"/>{col}<w:sz w:val="{int(size * 2)}"/></w:rPr>'
            f'<m:t xml:space="preserve">{t}</m:t></m:r>')


def omml_frac(num, den, color, size):
    n = ''.join(omml_run(t, it, color, size) for t, it in num) or omml_run(' ', False, color, size)
    d = ''.join(omml_run(t, it, color, size) for t, it in den) or omml_run(' ', False, color, size)
    xml = (f'<m:oMath xmlns:m="{M_NS}" {nsdecls("w")}><m:f><m:fPr><m:ctrlPr><w:rPr><w:rFonts w:ascii="Cambria Math" w:hAnsi="Cambria Math"/></w:rPr></m:ctrlPr></m:fPr>'
           f'<m:num>{n}</m:num><m:den>{d}</m:den></m:f></m:oMath>')
    return parse_xml(xml)


class Inline(HTMLParser):
    """Writes a snippet of inline HTML (b, i, sup, sub, br, span.fr fractions, span.key) into a paragraph."""

    def __init__(self, para, base):
        super().__init__(convert_charrefs=True)
        self.p = para
        self.base = base
        self.stack = []
        self.frac = None
        self.part = None

    def style(self):
        s = dict(self.base)
        for _, d in self.stack:
            s.update(d)
        return s

    def handle_starttag(self, tag, attrs):
        cls = (dict(attrs).get('class') or '').split()
        if tag == 'br':
            self.p.add_run().add_break()
            return
        d = {}
        if tag in ('b', 'strong'):
            d = {'bold': True}
        elif tag in ('i', 'em'):
            d = {'italic': True}
        elif tag == 'sup':
            d = {'sup': True}
        elif tag == 'sub':
            d = {'sub': True}
        elif tag == 'span':
            if 'fr' in cls:
                self.frac = {'num': [], 'den': []}
                d = {'frac': True}
            elif 'fn' in cls:
                self.part = 'num'
            elif 'fd' in cls:
                self.part = 'den'
            elif 'key' in cls:
                d = {'color': RED, 'bold': True}
        self.stack.append((tag, d))

    def handle_startendtag(self, tag, attrs):
        if tag == 'br':
            self.p.add_run().add_break()

    def handle_endtag(self, tag):
        while self.stack:
            t, d = self.stack.pop()
            if d.get('frac'):
                s = self.style()
                self.p._p.append(omml_frac(self.frac['num'], self.frac['den'], s.get('color'), s.get('size', 11)))
                self.frac = None
                self.part = None
            if t == tag:
                break

    def handle_data(self, data):
        if not data:
            return
        s = self.style()
        if self.frac is not None and self.part:
            txt = data.translate(SUPS) if s.get('sup') else data
            self.frac[self.part].append((txt, bool(s.get('italic'))))
            return
        r = self.p.add_run(data)
        font(r, s.get('font', SERIF), s.get('size', 11), s.get('bold', False), s.get('italic', False), s.get('color'), s.get('sup', False), s.get('sub', False))


def inline(p, html, **base):
    base.setdefault('font', SERIF)
    base.setdefault('size', 11)
    Inline(p, base).feed(str(html))


def plain_len(html):
    return len(re.sub(r'<[^>]+>', '', str(html)))


# ---------------------------------------------------------------- the document builder
class Builder:
    def __init__(self, data, guide, build_dir):
        self.d = data
        self.guide = guide
        self.build = build_dir
        self.doc = Document()
        self.setup()

    # ---------- document defaults, page, footer ----------
    def setup(self):
        doc = self.doc
        st = doc.styles['Normal']
        st.font.name = SERIF
        st.font.size = Pt(11)
        rpr = st.element.get_or_add_rPr()
        rf = rpr.find(qn('w:rFonts'))
        for k in ('w:ascii', 'w:hAnsi', 'w:cs', 'w:eastAsia'):
            rf.set(qn(k), SERIF)
        st.paragraph_format.space_before = Pt(0)
        st.paragraph_format.space_after = Pt(0)
        st.paragraph_format.line_spacing = 1.0
        s = doc.sections[0]
        self.page_setup(s)
        f = s.footer
        p = f.paragraphs[0]
        para_fmt(p, align='l', left=25.4 - M_LEFT)
        fld = el('w:fldSimple', **{'w:instr': 'PAGE'})
        r = el('w:r')
        rPr = el('w:rPr')
        rPr.append(el('w:rFonts', **{'w:ascii': PNFONT, 'w:hAnsi': PNFONT, 'w:cs': PNFONT}))
        rPr.append(el('w:sz', **{'w:val': 22}))
        r.append(rPr)
        t = el('w:t')
        t.text = '1'
        r.append(t)
        fld.append(r)
        p._p.append(fld)
        # Word 2013+ layout (not "Compatibility Mode"): table indents are measured to the table's edge
        for cs in doc.settings.element.iter(qn('w:compatSetting')):
            if cs.get(qn('w:name')) == 'compatibilityMode':
                cs.set(qn('w:val'), '15')
        doc.core_properties.title = self.title()
        doc.core_properties.author = 'Kingscliff High School'

    def page_setup(self, s):
        s.page_width, s.page_height = Mm(PAGE_W), Mm(PAGE_H)
        s.left_margin, s.right_margin = Mm(M_LEFT), Mm(M_RIGHT)
        s.top_margin, s.bottom_margin = Mm(M_TOP), Mm(M_BOTTOM)
        s.footer_distance = Mm(15.6)
        s.header_distance = Mm(13.5)

    def title(self):
        m = self.d['meta']
        t = f"{m['subject']} {m['course']} {m['examName']} {m['year']}"
        return t + (' Marking Guide' if self.guide else '')

    def fig_path(self, fid):
        return os.path.join(self.build, self.d['figures'][fid]['png'])

    def fig_w(self, fid, cap=None):
        w = self.d['figures'][fid]['wmm']
        return min(w, cap) if cap else w

    # ---------- front page ----------
    def front(self):
        doc, d, guide = self.doc, self.d, self.guide
        p = doc.paragraphs[0] if doc.paragraphs else doc.add_paragraph()
        para_fmt(p, line_exact=45.0 - M_TOP)
        # logo + NAME / TEACHER boxes: a borderless layout table, the boxes in a table nested in the right cell
        t = doc.add_table(rows=1, cols=3)
        table_setup(t, [62.7, 12.1, 72.1], borders=False, pad=(0, 0), indent=33.4 - M_LEFT)
        row_props(t.rows[0], 37.2, rule='exact', cant_split=False)
        logo, _, right = t.rows[0].cells
        lp = logo.paragraphs[0]
        para_fmt(lp)
        picture(lp.add_run(), d['logo'], 62.7)
        para_fmt(right.paragraphs[0], line_exact=0.35)
        bt = right.add_table(rows=4, cols=1)
        table_setup(bt, [72.1], borders=False, pad=(0, 0))
        for i, h in enumerate([3.9 - 0.35, 14 if guide else 11.3, 10.7, 11.3]):
            row_props(bt.rows[i], h, rule='exact', cant_split=False)
        box = (4, '444444')
        if guide:
            c = bt.cell(1, 0)
            set_cell_borders(c, top=(8, RED), left=(8, RED), bottom=(8, RED), right=(8, RED))
            c.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            pp = c.paragraphs[0]
            para_fmt(pp, align='c')
            font(pp.add_run('TEACHER COPY'), SANS, 12, True, color=RED)
            pp2 = c.add_paragraph()
            para_fmt(pp2, align='c')
            font(pp2.add_run('Not for students'), SANS, 9, color=RED)
        else:
            for ri, label in ((1, 'NAME:'), (3, 'TEACHER:')):
                c = bt.cell(ri, 0)
                set_cell_borders(c, top=box, left=box, bottom=box, right=box)
                set_cell_margins(c, top=3.1, left=2.2)
                pp = c.paragraphs[0]
                para_fmt(pp)
                font(pp.add_run(label), TIMES, 11)
        para_fmt(right.paragraphs[-1], line_exact=0.35)
        # title lines
        lines = list(d['meta']['titleLines'])
        pitch = 13.4 if guide else 14.7
        top = 86.0 if guide else 89.0
        p = doc.add_paragraph()
        para_fmt(p, line_exact=top - 82.2 - 1.4)
        for i, line in enumerate(lines + (['Marking Guide'] if guide else [])):
            p = doc.add_paragraph()
            para_fmt(p, align='c', line_exact=pitch)
            mg = guide and i == len(lines)
            font(p.add_run(line), SANS, 22 if mg else 25, bold=mg, color=RED if mg else NAVY)
        # Instructions | bar | Sections and Marks
        box_top = 156.0 if guide else 153.8
        used = top + pitch * (len(lines) + (1 if guide else 0))
        p = doc.add_paragraph()
        para_fmt(p, line_exact=max(0.4, box_top - used + (1.5 if guide else 1.0)))
        t = doc.add_table(rows=1, cols=3)
        table_setup(t, [74.8, 5.3, 75.0], borders=False, pad=(3.0, 2.6), indent=27.45 - M_LEFT)
        row_props(t.rows[0], 73 if guide else 84.7, rule='exact', cant_split=False)
        edge = (8, '000000')
        L, B, R = t.rows[0].cells
        set_cell_borders(L, top=edge, left=edge, bottom=edge)
        set_cell_borders(B, top=edge, bottom=edge)
        set_cell_borders(R, top=edge, right=edge, bottom=edge)
        set_cell_shading(B, KBLUE)
        set_cell_margins(B, left=0, right=0)
        hp = L.paragraphs[0]
        para_fmt(hp, line=1.0)
        font(hp.add_run('Notes for markers' if guide else 'Instructions'), SANS, 12, True, color=NAVY)
        items = d['markingNotes'] if guide else d['instructions']
        size = 10 if guide else 12
        for i, it in enumerate(items):
            p = L.add_paragraph()
            para_fmt(p, before=(4 if guide else 8.6) if i == 0 else 0, after=2 if guide else 4.85, line=1.0,
                     left=(7 if guide else 12.7) + 1.6, hanging=6.2, tabs=[(7 if guide else 12.7) + 1.6])
            font(p.add_run(f'{i + 1}.\t'), SANS, size)
            inline(p, it, font=SANS, size=size)
        hp = R.paragraphs[0]
        para_fmt(hp, line=1.0)
        font(hp.add_run('Sections and Marks'), SANS, 12, True, color=NAVY)
        # the sections list: label | mark, the mark on the label's last line (as in the PDF)
        lw = 75.0 - 2 * 2.6 - 19
        sp = R.add_paragraph()
        para_fmt(sp, line_exact=4 if guide else 6.2)
        st = R.add_table(rows=0, cols=2)
        table_setup(st, [lw, 19], borders=False, pad=(0, 0))
        rows = [(S['frontLabel'], f"/{S['marks']}", f"{S['green']} green · {S['blue']} blue" if guide else None, False) for S in d['sections']]
        rows.append(('Total marks', f"/{d['total']}", None, True))
        for i, (lab, mk, sub, tot) in enumerate(rows):
            row = st.add_row()
            row_props(row, cant_split=False)
            a, b2 = row.cells
            a.width, b2.width = Mm(lw), Mm(19)
            a.vertical_alignment = b2.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.BOTTOM
            gap = (3 if guide else 5.4) if not tot else 0
            pa = a.paragraphs[0]
            para_fmt(pa, line=1.0, before=1.5 if tot else 0, after=0 if sub else gap)
            font(pa.add_run(lab), SANS, 12 if tot else 11, tot)
            if sub:
                ps = a.add_paragraph()
                para_fmt(ps, before=0.6, after=gap)
                font(ps.add_run(sub), SANS, 8.5, color='333333')
            pb = b2.paragraphs[0]
            para_fmt(pb, line=1.0, after=gap)
            font(pb.add_run(mk), SANS, 12 if tot else 11, tot)
        para_fmt(R.paragraphs[-1], line_exact=0.5)
        if guide:
            self.distribution()

    def distribution(self):
        doc, d = self.doc, self.d
        p = doc.add_paragraph()
        para_fmt(p, line_exact=6)
        widths = [52, 15, 15, 20, 20, 33.1]
        t = doc.add_table(rows=0, cols=6)
        table_setup(t, widths, pad=(0.9, 1.6), indent=27.45 - M_LEFT)
        total = d['total'] or 1
        pct = lambda x: f"{round(x / total * 1000) / 10:g}%"
        rows = [['Section', 'Marks', 'Share', ('g', 'Green'), ('b', 'Blue'), 'Questions']]
        for S in d['sections']:
            g = S['groups']
            rows.append([S['frontLabel'], str(S['marks']), pct(S['marks']), str(S['green']), str(S['blue']), f"{g['mc']['n']} MC · {g['sa']['n']} SA · {g['worded']['n']} WP"])
        rows.append(['Total', str(d['total']), '100%', f"{d['green']} ({pct(d['green'])})", f"{d['blue']} ({pct(d['blue'])})", ''])
        for ri, r in enumerate(rows):
            row = t.add_row()
            row_props(row, cant_split=True)
            for ci, v in enumerate(r):
                c = row.cells[ci]
                c.width = Mm(widths[ci])
                p = c.paragraphs[0]
                para_fmt(p, align='l' if ci in (0, 5) else 'c')
                bold = ri == 0 or ri == len(rows) - 1
                if isinstance(v, tuple):
                    font(p.add_run('■ ' if v[0] == 'g' else '● '), SANS, 9, color=GREEN if v[0] == 'g' else KBLUE)
                    v = v[1]
                font(p.add_run(v), SANS, 9, bold)
                if ri == 0:
                    set_cell_shading(c, TINT)
                elif ri == len(rows) - 1:
                    set_cell_shading(c, TINT2)

    # ---------- blocks inside cells ----------
    def blocks(self, cell, blocks, used, width, align='l', size=11, keep=False):
        for b in blocks or []:
            t = b['t']
            if t == 'p':
                cls = (b.get('cls') or '').split()
                p = next_para(cell, used)
                al = 'c' if 'c' in cls else align
                if 'part' in cls:
                    para_fmt(p, align='l', before=0.6, after=0.6, left=6.0, hanging=5.5, tabs=[6.0], keep_next=keep)
                    font(p.add_run(f"{b.get('label', '')}.\t"), SERIF, size)
                    inline(p, b['html'], size=size)
                elif 'key' in cls:
                    para_fmt(p, align='l', after=0.5, before=0.8 if 'ans' in cls else 0, keep_next=keep)
                    inline(p, b['html'], size=10, color=RED, bold='ans' in cls)
                else:
                    para_fmt(p, align=al, before=1.6 if 'formula' in cls else 0, after=1.2, keep_next=keep)
                    inline(p, b['html'], size=size)
            elif t == 'fig':
                p = next_para(cell, used)
                para_fmt(p, align='c', before=1.6, after=1.0, keep_next=keep)
                picture(p.add_run(), self.fig_path(b['id']), self.fig_w(b['id'], width))
            elif t == 'table':
                self.data_table(cell, b, used, width)
            elif t == 'figs':
                self.fig_row(cell, b['items'], used, width)

    def data_table(self, cell, b, used, width):
        rows = b['rows']
        ncol = max(len(r) for r in rows)
        if b.get('colW'):
            widths = list(b['colW'])
        else:
            widths = []
            for ci in range(ncol):
                longest = max(plain_len(r[ci].get('text', '') if isinstance(r[ci], dict) and r[ci].get('key') is None else r[ci].get('key', '')) for r in rows if ci < len(r))
                widths.append(max(10.0, min(42.0, longest * 2.0 + 5.4)))
            if sum(widths) > width - 2:
                k = (width - 2) / sum(widths)
                widths = [w * k for w in widths]
        if not used[0]:
            used[0] = True
            para_fmt(cell.paragraphs[0], line_exact=1.2)
        t = cell.add_table(rows=0, cols=ncol)
        indent = max(0.0, (width - sum(widths)) / 2)
        table_setup(t, widths, pad=(0.7, 2.4), indent=indent, sz=5)
        tblPr = t._tbl.tblPr
        for ri, r in enumerate(rows):
            row = t.add_row()
            row_props(row)
            for ci in range(ncol):
                c = row.cells[ci]
                c.width = Mm(widths[ci])
                c.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
                v = r[ci] if ci < len(r) else {'text': ''}
                head = (b.get('headRow') and ri == 0) or (b.get('headCol') and ci == 0)
                foot = b.get('footRow') and ri == len(rows) - 1
                p = c.paragraphs[0]
                para_fmt(p, align='l' if b.get('align') == 'left' else 'c')
                if head:
                    set_cell_shading(c, TINT2)
                if v.get('key') is not None:
                    if self.guide:
                        inline(p, v['key'], size=10.5, color=RED, bold=True)
                else:
                    inline(p, v.get('text', ''), size=10.5, bold=bool(head or foot or v.get('b')))
        # the paragraph Word needs after a table in a cell: keep it tiny
        last = cell.paragraphs[-1]
        para_fmt(last, line_exact=1.2)
        del tblPr

    def fig_row(self, cell, items, used, width):
        if not used[0]:
            used[0] = True
            para_fmt(cell.paragraphs[0], line_exact=1.2)
        n = len(items)
        t = cell.add_table(rows=1, cols=n)
        widths = [width / n] * n
        table_setup(t, widths, borders=False, pad=(0, 0.5))
        for i, it in enumerate(items):
            c = t.rows[0].cells[i]
            c.width = Mm(widths[i])
            c.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            u = [False]
            if it['t'] == 'fig':
                p = next_para(c, u)
                para_fmt(p, align='c')
                picture(p.add_run(), self.fig_path(it['id']), self.fig_w(it['id'], widths[i] - 1))
            elif it['t'] == 'table':
                self.data_table(c, it, u, widths[i] - 1)
            else:
                p = next_para(c, u)
                para_fmt(p, align='c')
                inline(p, it.get('html', ''))
        para_fmt(cell.paragraphs[-1], line_exact=1.2)

    def mark_text(self, cell, text, size=11, bold=False, align='c'):
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        if cell.width is not None and cell.width < Mm(13):
            set_cell_margins(cell, left=0.4, right=0.4)
        p = cell.paragraphs[0]
        para_fmt(p, align=align)
        if text != '':
            font(p.add_run(str(text)), SERIF, size, bold)

    def keep_row(self, row):
        for c in row.cells:
            for p in c.paragraphs:
                p.paragraph_format.keep_with_next = True

    # ---------- section tables ----------
    def section(self, S, first):
        doc = self.doc
        if first is False:
            tiny_para(doc, page_break=True)
        t = doc.add_table(rows=0, cols=4)
        table_setup(t, COLS)
        for r in S['rows']:
            self.row(t, r)
        return t

    def row(self, t, r):
        k = r['k']
        row = t.add_row()
        for i, c in enumerate(row.cells):
            c.width = Mm(COLS[i])
        num = row.cells[0]
        num.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        set_cell_margins(num, left=0.3, right=0.3)
        if k == 'title':
            row_props(row, 11.7, header=True)
            c = merge(row, 1, 3)
            for x in (num, c):
                set_cell_shading(x, NAVY)
            self.mark_text(c, r['text'], 16, True)
            c.paragraphs[0].runs[0].font.color.rgb = RGBColor.from_string('FFFFFF')
        elif k == 'sub':
            row_props(row, 9.1)
            if r.get('markHead'):
                c = merge(row, 1, 2)
                self.mark_text(c, r['text'], 11, True)
                m = row.cells[3]
                set_cell_margins(m, left=0.5, right=0.5)
                self.mark_text(m, r['markHead'], 9, True)
                cells = [num, c, m]
            else:
                c = merge(row, 1, 3)
                self.mark_text(c, r['text'], 11, True)
                cells = [num, c]
            for x in cells:
                set_cell_shading(x, TINT)
        elif k == 'colhdr':
            row_props(row, 9.1)
            for i, txt in enumerate(['', 'Question', 'Working Out', 'Mark']):
                self.mark_text(row.cells[i], txt, 11, True)
                set_cell_shading(row.cells[i], TINT2)
        elif k == 'total':
            row_props(row, 9.1)
            c = merge(row, 1, 2)
            self.mark_text(c, 'Total', 11, True)
            m = row.cells[3]
            self.mark_text(m, f"/{r['marks']}", 11, True)
            for x in (num, c, m):
                set_cell_shading(x, TINT2)
        elif k == 'info':
            row_props(row)
            c = merge(row, 1, 2)
            self.blocks(c, r['q'], [False], COLS[1] + COLS[2] - 2 * CELL_PAD_X, align='c')
        elif k == 'mc':
            row_props(row)
            self.mark_text(num, r['n'])
            c = merge(row, 1, 2)
            set_cell_margins(c, top=2.6, bottom=3.0, left=3.0, right=3.0)
            self.mark_text(row.cells[3], r['marks'])
            self.mc_cell(c, r)
        elif k == 'q':
            parts = r['parts']
            total_h = sum(p['space'] for p in parts)
            row_props(row, total_h)
            self.mark_text(num, r['n'])
            q, wk, mk = row.cells[1], row.cells[2], row.cells[3]
            set_cell_margins(mk, left=0.4, right=0.4)
            self.blocks(q, r['q'], [False], COLS[1] - 2 * CELL_PAD_X)
            self.work_cell(wk, parts)
            if len(parts) > 1:
                self.marks_stack(mk, parts)
            else:
                self.mark_text(mk, parts[0]['marks'])
        elif k == 'stem':
            row_props(row)
            self.mark_text(num, r['n'])
            c = merge(row, 1, 2)
            self.mark_text(row.cells[3], '')
            width = COLS[1] + COLS[2] - 2 * CELL_PAD_X
            if r.get('figPos') == 'right' and r.get('fig'):
                fw = max(self.fig_w(f['id']) for f in r['fig'] if f['t'] == 'fig') + 2 if any(f['t'] == 'fig' for f in r['fig']) else 70
                fw = min(fw, width * 0.6)
                c.paragraphs[0].paragraph_format.line_spacing_rule = WD_LINE_SPACING.EXACTLY
                c.paragraphs[0].paragraph_format.line_spacing = Mm(0.6)
                nt = c.add_table(rows=1, cols=2)
                table_setup(nt, [width - fw, fw], borders=False, pad=(0, 0))
                a, bcell = nt.rows[0].cells
                a.width, bcell.width = Mm(width - fw), Mm(fw)
                self.blocks(a, r['q'] + (r.get('formula') or []), [False], width - fw - 4)
                set_cell_margins(a, right=4)
                self.blocks(bcell, r['fig'], [False], fw)
                para_fmt(c.paragraphs[-1], line_exact=0.6)
            else:
                self.blocks(c, r['q'] + (r.get('fig') or []) + (r.get('formula') or []), [False], width)
        elif k == 'part':
            row_props(row, r['space'])
            self.mark_text(num, r['label'])
            self.blocks(row.cells[1], r['q'], [False], COLS[1] - 2 * CELL_PAD_X)
            self.work_cell(row.cells[2], [{'label': '', 'space': r['space'], 'key': r.get('key', [])}])
            self.mark_text(row.cells[3], r['marks'])
        if r.get('kn'):
            self.keep_row(row)

    def work_cell(self, cell, parts):
        """Part labels a., b. at the top of each part's block; in the marking guide, the worked answer in red."""
        used = [False]
        prev_used = 0.0
        prev_h = 0.0
        for i, p in enumerate(parts):
            before = 0 if i == 0 else max(0.0, prev_h - prev_used)
            this_used = 0.0
            if p.get('label'):
                para = next_para(cell, used)
                para_fmt(para, before=before)
                font(para.add_run(f"{p['label']}."), SERIF, 11)
                this_used += L11
                before = 0
            if self.guide:
                for j, b in enumerate(p.get('key') or []):
                    para = next_para(cell, used)
                    ans = 'ans' in (b.get('cls') or '')
                    para_fmt(para, before=before + (0.8 if ans else (0.4 if j == 0 and p.get('label') else 0)), after=0.5)
                    inline(para, b['html'], size=10, color=RED, bold=ans)
                    lines = max(1, -(-plain_len(b['html']) // 52))
                    this_used += lines * L10 + 0.5 + (0.8 if ans else 0)
                    before = 0
            if this_used == 0 and i > 0:
                para = next_para(cell, used)
                para_fmt(para, before=before)
            prev_used = this_used + (2 * CELL_PAD_Y if i == 0 else 0)
            prev_h = p['space']

    def marks_stack(self, cell, parts):
        """Marks for inline parts, each centred on its part's working block."""
        used = [False]
        for i, p in enumerate(parts):
            if i == 0:
                before = p['space'] / 2 - L11 / 2 - CELL_PAD_Y
            else:
                before = parts[i - 1]['space'] / 2 + p['space'] / 2 - L11
            para = next_para(cell, used)
            para_fmt(para, align='c', before=before)
            font(para.add_run(str(p['marks'])), SERIF, 11)

    def mc_cell(self, c, r):
        width = COLS[1] + COLS[2] - 6.0
        if r.get('figPos') == 'right' and r.get('fig'):
            fw = max([self.fig_w(f['id']) for f in r['fig'] if f['t'] == 'fig'] or [60]) + 2
            fw = min(fw, width * 0.55)
            para_fmt(c.paragraphs[0], line_exact=0.6)
            nt = c.add_table(rows=1, cols=2)
            table_setup(nt, [width - fw, fw], borders=False, pad=(0, 0))
            a, b = nt.rows[0].cells
            a.width, b.width = Mm(width - fw), Mm(fw)
            a.vertical_alignment = b.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(a, right=4)
            self.blocks(a, r['stem'], [False], width - fw - 4, align='c')
            self.options(a, r, width - fw - 4)
            self.blocks(b, r['fig'], [False], fw)
            para_fmt(c.paragraphs[-1], line_exact=0.6)
        else:
            used = [False]
            self.blocks(c, r['stem'], used, width, align='c')
            self.blocks(c, r.get('fig') or [], used, width)
            self.options(c, r, width)

    def options(self, cell, r, width):
        lay = r['optLayout']
        opts = r['options']
        if lay == 'stack':
            for i, o in enumerate(opts):
                p = cell.add_paragraph()
                para_fmt(p, align='l', before=3.6 if i == 0 else 0, after=1.8, left=18 + 6.2, hanging=6.2, tabs=[18 + 6.2])
                self.option_runs(p, i, o, r, width)
            return
        ncol = 4 if lay == 'row' else 2
        nrow = 4 // ncol
        lead = 9.0 if lay in ('row', 'grid') else 0
        colw = (width - lead - 4) / ncol
        p0 = cell.add_paragraph()
        para_fmt(p0, line_exact=3.6)
        t = cell.add_table(rows=nrow, cols=ncol)
        table_setup(t, [colw] * ncol, borders=False, pad=(0, 0), indent=lead)
        for i, o in enumerate(opts):
            c = t.cell(i // ncol, i % ncol)
            c.width = Mm(colw)
            p = c.paragraphs[0]
            para_fmt(p, align='l', after=1.8 if i // ncol < nrow - 1 else 0, left=6.2, hanging=6.2, tabs=[6.2])
            self.option_runs(p, i, o, r, colw - 7)
        para_fmt(cell.paragraphs[-1], line_exact=0.6)

    def option_runs(self, p, i, o, r, width):
        right = self.guide and i == r['answer']
        lr = p.add_run(f'{LETTERS[i]}.')
        font(lr, SERIF, 11, bold=right, color=RED if right else None)
        if right:
            rPr = lr._r.get_or_add_rPr()
            rPr.append(el('w:bdr', **{'w:val': 'single', 'w:sz': 6, 'w:space': 0, 'w:color': RED}))
        font(p.add_run('\t'), SERIF, 11)
        if o.get('fig'):
            picture(p.add_run(), self.fig_path(o['fig']['id']), self.fig_w(o['fig']['id'], width))
        else:
            inline(p, o['html'], size=11, bold=right, color=RED if right else None)

    # ---------- marking guide pages ----------
    def summary(self):
        doc, d = self.doc, self.d
        widths = [13, 71, 14, 20, 66]
        for si, s in enumerate(d['summary']):
            if si == 0:
                tiny_para(doc, page_break=True)
            else:
                p = doc.add_paragraph()
                para_fmt(p, line_exact=5)
            t = doc.add_table(rows=0, cols=5)
            table_setup(t, widths, pad=(1.0, 1.6))
            row = t.add_row()
            row_props(row, 9.5, header=True)
            c = merge(row, 0, 4, widths)
            set_cell_shading(c, NAVY)
            self.mark_text(c, f"{s['heading']}: answer summary", 13, True)
            c.paragraphs[0].runs[0].font.color.rgb = RGBColor.from_string('FFFFFF')
            self.keep_row(row)
            row = t.add_row()
            row_props(row, 7, header=True)
            for i, h in enumerate(['Q', 'Answer', 'Marks', 'Level', 'Topic and source']):
                row.cells[i].width = Mm(widths[i])
                self.mark_text(row.cells[i], h, 10, True)
                set_cell_shading(row.cells[i], TINT2)
            self.keep_row(row)
            for ri, r in enumerate(s['rows']):
                row = t.add_row()
                row_props(row)
                for i in range(5):
                    row.cells[i].width = Mm(widths[i])
                self.mark_text(row.cells[0], r['q'], 10, True)
                row.cells[0].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
                p = row.cells[1].paragraphs[0]
                para_fmt(p)
                inline(p, r['answer'], size=10)
                self.mark_text(row.cells[2], r['marks'], 10)
                row.cells[2].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
                p = row.cells[3].paragraphs[0]
                para_fmt(p)
                self.level_runs(p, r['level'], 10)
                p = row.cells[4].paragraphs[0]
                para_fmt(p)
                if r.get('topic'):
                    inline(p, r['topic'], size=10)
                    p = row.cells[4].add_paragraph()
                    para_fmt(p)
                inline(p, r.get('src', ''), size=9, color='444444')
                if ri == len(s['rows']) - 1:
                    self.keep_row(row)
            row = t.add_row()
            row_props(row, 7)
            for i in range(5):
                row.cells[i].width = Mm(widths[i])
            self.mark_text(row.cells[1], 'Total', 10, True)
            self.mark_text(row.cells[2], s['marks'], 10, True)
            c = merge(row, 3, 4, widths)
            self.mark_text(c, f"{s['green']} green · {s['blue']} blue", 10, False, 'l')
            for x in row.cells:
                set_cell_shading(x, TINT2)

    def level_runs(self, p, level, size):
        if level == 'green':
            font(p.add_run('■ '), SERIF, size, color=GREEN)
            font(p.add_run('Green'), SERIF, size)
        elif level == 'blue':
            font(p.add_run('● '), SERIF, size, color=KBLUE)
            font(p.add_run('Blue'), SERIF, size)
        else:
            font(p.add_run(str(level or '')), SERIF, size)

    def criteria(self):
        doc, d = self.doc, self.d
        tiny_para(doc, page_break=True)
        for s in d['criteria']:
            p = doc.add_paragraph()
            para_fmt(p, align='c', before=1, after=2.5, keep_next=True)
            put(p._p.get_or_add_pPr(), el('w:shd', **{'w:val': 'clear', 'w:color': 'auto', 'w:fill': NAVY}), PPR)
            font(p.add_run(f"{s['heading']}: marking criteria"), SERIF, 13, True, color='FFFFFF')
            for q in s['questions']:
                self.crit_question(q)

    def crit_question(self, q):
        doc = self.doc
        widths = [9, 175]
        t = doc.add_table(rows=0, cols=2)
        table_setup(t, widths, pad=(1.2, 2.0))
        row = t.add_row()
        row_props(row)
        c = merge(row, 0, 1, widths)
        set_cell_shading(c, TINT)
        p = c.paragraphs[0]
        para_fmt(p, tabs=[30, 48])
        font(p.add_run(f"Question {q['n']}"), SERIF, 10.5, True)
        font(p.add_run(f"\t{q['marks']} mark{'s' if q['marks'] != 1 else ''}\t"), SERIF, 10)
        for i, lv in enumerate(q['levels']):
            if i:
                font(p.add_run(' + '), SERIF, 10)
            self.level_runs(p, lv, 10)
        font(p.add_run('    '), SERIF, 10)
        inline(p, q.get('topic', ''), size=10)
        p2 = c.add_paragraph()
        para_fmt(p2, align='r')
        inline(p2, q.get('src', ''), size=9, color='333333')
        self.keep_row(row)
        multi = len(q['parts']) > 1
        for pi, part in enumerate(q['parts']):
            row = t.add_row()
            row_props(row)
            row.cells[0].width, row.cells[1].width = Mm(widths[0]), Mm(widths[1])
            lp = row.cells[0].paragraphs[0]
            para_fmt(lp)
            if part.get('label'):
                font(lp.add_run(f"{part['label']})"), SERIF, 10, True)
            cell = row.cells[1]
            used = [False]

            def line(label, html=None, **kw):
                para = next_para(cell, used)
                para_fmt(para, left=17, hanging=17, tabs=[17], after=0.8)
                font(para.add_run(f'{label}\t' if label else '\t'), SERIF, 8.5, color='444444')
                if html is not None:
                    inline(para, html, size=10, **kw)
                return para

            para = line('ANSWER', part['answer'], color=RED, bold=True)
            if multi:
                font(para.add_run(f"     ({part['marks']} mark{'s' if part['marks'] != 1 else ''}, "), SERIF, 10, color='333333')
                self.level_runs(para, part['level'], 10)
                font(para.add_run(')'), SERIF, 10, color='333333')
            for i, w in enumerate(part.get('working') or []):
                line('WORKING' if i == 0 else '', w, color=RED)
            for i, (m, txt) in enumerate(part.get('marking') or []):
                para = line('MARKS' if i == 0 else '')
                font(para.add_run(f"{m} mark{'s' if m != 1 else ''}: "), SERIF, 10, True)
                inline(para, txt, size=10)
            for i, n in enumerate(part.get('notes') or []):
                line('NOTES' if i == 0 else '', n, italic=True)
            if pi < len(q['parts']) - 1:
                self.keep_row(row)
        p = doc.add_paragraph()
        para_fmt(p, line_exact=2.5)

    def guide_header(self, s):
        s.header.is_linked_to_previous = False
        p = s.header.paragraphs[0]
        para_fmt(p, align='r')
        r = p.add_run(' Marking guide: answers in red ')
        font(r, SANS, 9, True, color=RED)
        r._r.get_or_add_rPr().append(el('w:bdr', **{'w:val': 'single', 'w:sz': 6, 'w:space': 1, 'w:color': RED}))
        s.footer.is_linked_to_previous = True

    # ---------- whole documents ----------
    def build_doc(self):
        self.front()
        if self.guide:
            self.summary()
            self.criteria()
            sec = self.doc.add_section(WD_SECTION.NEW_PAGE)
            self.page_setup(sec)
            self.guide_header(sec)
            for i, S in enumerate(self.d['sections']):
                self.section(S, first=(i == 0))
        else:
            for S in self.d['sections']:
                self.section(S, first=False)
        return self.doc


def main():
    ap = argparse.ArgumentParser()
    here = os.path.dirname(os.path.abspath(__file__))
    ap.add_argument('--json', default=os.path.join(here, 'build', 'exam.json'))
    ap.add_argument('--out', default=os.path.join(here, '..', 'Yearly exam 2026'))
    a = ap.parse_args()
    with open(a.json, encoding='utf-8') as f:
        data = json.load(f)
    build_dir = os.path.dirname(os.path.abspath(a.json))
    os.makedirs(a.out, exist_ok=True)
    for guide, name in ((False, data['files']['exam']), (True, data['files']['guide'])):
        doc = Builder(data, guide, build_dir).build_doc()
        path = os.path.join(a.out, f'{name}.docx')
        doc.save(path)
        print(f'  {name}.docx')


if __name__ == '__main__':
    main()
